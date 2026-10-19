"""app/services/export_service.py — 审查报告导出（Word）

PRD P1：合同审查结果导出为 Word / PDF。
这里实现 Word（.docx）—— 服务器无需安装 Office/LibreOffice，python-docx 直出；
需要 PDF 时前端用浏览器打印即可（版式与本文件一致）。
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

LEVEL_LABEL = {"high": "高风险", "medium": "中风险", "low": "低风险"}
OVERALL_LABEL = {"high": "高风险", "medium": "中风险", "low": "低风险"}
STATUS_LABEL = {
    "pending": "待处理",
    "confirmed": "已确认",
    "ignored": "已忽略",
    "todo": "待修改",
}


def _set_base_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def build_review_docx(review, items: list, contract) -> bytes:
    """生成审查报告 .docx，返回字节流"""
    doc = Document()
    _set_base_font(doc)

    # —— 标题 ——
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("合同审查报告")
    run.font.size = Pt(20)
    run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"合同文件：{contract.filename}").font.size = Pt(10)

    # —— 基本信息表 ——
    doc.add_paragraph()
    overall = OVERALL_LABEL.get(
        review.overall_risk.value if review.overall_risk else "", "未评定"
    )
    risk_items = [i for i in items if not i.is_positive]
    positives = [i for i in items if i.is_positive]
    high = sum(1 for i in risk_items if i.risk_level.value == "high")
    medium = sum(1 for i in risk_items if i.risk_level.value == "medium")
    low = sum(1 for i in risk_items if i.risk_level.value == "low")

    table = doc.add_table(rows=5, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("合同名称", contract.filename),
        ("审查时间", review.generated_at.strftime("%Y-%m-%d %H:%M")),
        ("总体风险", overall),
        ("风险项统计", f"高 {high} 条 / 中 {medium} 条 / 低 {low} 条"),
        ("正向发现", f"{len(positives)} 条"),
    ]
    for idx, (key, value) in enumerate(rows):
        table.cell(idx, 0).text = key
        table.cell(idx, 1).text = str(value)

    # —— 风险项明细 ——
    doc.add_paragraph()
    heading = doc.add_paragraph()
    heading.add_run("一、风险项明细").bold = True

    if not risk_items:
        doc.add_paragraph("本次审查未发现风险项。")
    for idx, item in enumerate(risk_items, start=1):
        level = LEVEL_LABEL.get(item.risk_level.value, item.risk_level.value)
        status = STATUS_LABEL.get(
            getattr(item, "status", None).value if getattr(item, "status", None) else "pending",
            "待处理",
        )
        p = doc.add_paragraph()
        r = p.add_run(f"{idx}. [{level}] {item.category_name}")
        r.bold = True
        if item.risk_level.value == "high":
            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

        doc.add_paragraph(f"风险等级：{level}　处置状态：{status}")
        if item.contract_snippet:
            doc.add_paragraph(f"合同原文：{item.contract_snippet}")
        if item.issue:
            doc.add_paragraph(f"问题说明：{item.issue}")
        if item.suggestion:
            doc.add_paragraph(f"修改建议：{item.suggestion}")

    # —— 正向发现 ——
    if positives:
        doc.add_paragraph()
        h2 = doc.add_paragraph()
        h2.add_run("二、正向发现").bold = True
        for item in positives:
            doc.add_paragraph(f"✓ [{item.category_name}] {item.issue or item.contract_snippet}")

    # —— 免责声明 ——
    doc.add_paragraph()
    note = doc.add_paragraph()
    nrun = note.add_run(
        "本报告由合同智能审查系统基于企业知识库自动生成，仅供参考，不构成法律意见。"
        "重大合同请交由执业律师复核。"
    )
    nrun.font.size = Pt(9)
    nrun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_paragraph(
        f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ).runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
