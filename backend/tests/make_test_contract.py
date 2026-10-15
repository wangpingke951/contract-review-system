"""生成一份测试合同.docx，用于测试 Phase 4 多模态解析"""
import sys as _sys
_sys.path.insert(0, __import__("os").path.dirname(__import__("os").path.abspath(__file__)))
from helpers import _asset
from docx import Document

doc = Document()

doc.add_heading('采购合同', level=0)

doc.add_paragraph('当事人（甲方）：北京科技有限公司')
doc.add_paragraph('当事人（乙方）：上海贸易有限公司')

doc.add_heading('第一条 合同标的', level=1)
doc.add_paragraph('甲方向乙方采购计算机设备一批，具体型号、数量、单价见附件一《设备清单》。乙方应在合同签订后 15 个工作日内完成交货。')

doc.add_heading('第二条 合同金额', level=1)
doc.add_paragraph('合同总金额为人民币壹佰万元整（¥1,000,000.00），不含增值税。税费由甲方另行承担。')

doc.add_heading('第三条 付款方式', level=1)
doc.add_paragraph('乙方发货并提供足额增值税专用发票后 30 日内，甲方以银行转账方式支付全额货款。')

doc.add_heading('第四条 违约责任', level=1)
doc.add_paragraph('任何一方违反本合同约定，应向对方支付合同总金额 5% 的违约金。因一方违约给对方造成其他损失的，违约方应承担全部赔偿责任。')

doc.add_heading('第五条 争议解决', level=1)
doc.add_paragraph('因本合同产生的争议，双方应友好协商解决；协商不成的，提交北京市海淀区人民法院诉讼解决。')

doc.add_heading('第六条 保密条款', level=1)
doc.add_paragraph('双方应对合同内容及履约过程中获悉的对方商业秘密严格保密，未经对方书面同意不得向任何第三方披露。')

doc.save(_asset('test_contract.docx'))
print('✅ test_contract.docx 已生成')
