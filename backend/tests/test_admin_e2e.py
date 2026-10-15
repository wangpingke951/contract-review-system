"""Phase 8 端到端测试：知识库上传→入库→删除 + 风险类别增删 + 权限校验"""
import sys, asyncio
import httpx
sys.path.insert(0, __import__("os").path.join(__import__("os").path.dirname(__import__("os").path.abspath(__file__)), ".."))
from helpers import _asset

BASE = "http://localhost:8000"

KNOWLEDGE_TEXT = """审查指引：服务类合同风险要点

第一条 服务范围
服务类合同必须明确服务内容、交付标准与验收方式。仅约定"提供相关服务"而无数量化标准的，属于范围条款缺陷，应提示补充服务水平协议（SLA）。

第二条 服务费用
服务费建议采用"基础费用+绩效挂钩"结构。一次性预付全年的条款对采购方不利，建议按季度或半年度分期支付。

第三条 人员管理
应约定服务人员名单及关键岗位人员更换须经采购方书面同意，防止服务质量波动。人员社保与劳动关系责任应由服务方承担。

第四条 数据安全
涉及数据处理的服务合同应约定数据归属、保密义务、泄露责任及合同终止后的数据返还与销毁义务。
"""

CATEGORIES = []


def make_knowledge_docx(path):
    from docx import Document
    doc = Document()
    doc.add_heading("服务类合同审查指引", level=0)
    for para in KNOWLEDGE_TEXT.split("\n"):
        if para.strip():
            doc.add_paragraph(para)
    doc.save(path)


async def main():
    import os
    docx_path = _asset("knowledge_service.docx")
    make_knowledge_docx(docx_path)

    # 1) admin 登录
    r = httpx.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": "admin123"})
    admin_token = r.json()["access_token"]
    AH = {"Authorization": f"Bearer {admin_token}"}
    print("✅ admin 登录")

    # 2) 普通用户登录 → 访问 admin 接口应 403
    r = httpx.post(f"{BASE}/api/auth/login", json={"username": "user", "password": "user123"})
    user_token = r.json()["access_token"]
    r = httpx.get(f"{BASE}/api/admin/categories", headers={"Authorization": f"Bearer {user_token}"})
    assert r.status_code == 403, f"应 403, 实际 {r.status_code}"
    print("✅ 普通用户访问 admin 接口被拒 (403)")

    # 3) 上传知识文档（解析 → 切块 → bge-m3 向量化 → ChromaDB）
    print("⏳ 上传知识文档（向量化中...）")
    with open(docx_path, "rb") as f:
        r = httpx.post(f"{BASE}/api/admin/knowledge/upload",
                       files={"file": ("服务类合同审查指引.docx", f,
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                       data={"category": "guide"}, headers=AH, timeout=120)
    print(f"   HTTP {r.status_code}")
    if r.status_code != 200:
        print(r.text[:400]); sys.exit(1)
    doc = r.json()
    print(f"✅ 知识文档入库: id={doc['id']}, status={doc['status']}, chunks={doc['chunk_count']}")
    assert doc["status"] == "enabled" and doc["chunk_count"] >= 2

    # 4) 知识文档列表
    r = httpx.get(f"{BASE}/api/admin/knowledge", headers=AH)
    docs = r.json()
    print(f"✅ 知识列表: {len(docs)} 条")

    # 5) 风险类别：新增 → 重名拒绝 → 删除 → 全删光时拒绝
    r = httpx.post(f"{BASE}/api/admin/categories", json={"name": "数据安全"}, headers=AH)
    assert r.status_code == 200, r.text
    new_cat = r.json()
    print(f"✅ 新增类别: id={new_cat['id']}, name={new_cat['name']}")

    r = httpx.post(f"{BASE}/api/admin/categories", json={"name": "数据安全"}, headers=AH)
    assert r.status_code == 400
    print("✅ 重名类别被拒绝 (400)")

    r = httpx.delete(f"{BASE}/api/admin/categories/{new_cat['id']}", headers=AH)
    assert r.status_code == 200
    print("✅ 删除类别成功")

    # 只剩 1 个时删除应被拒
    r = httpx.get(f"{BASE}/api/admin/categories", headers=AH)
    cats = r.json()
    if len(cats) == 1:
        r = httpx.delete(f"{BASE}/api/admin/categories/{cats[0]['id']}", headers=AH)
        assert r.status_code == 400
        print("✅ 最后一个类别删除被拒 (400)")
    else:
        print(f"✅ 类别数 {len(cats)} > 1，跳过最后一条约束测试")

    # 6) 删除知识文档 → ChromaDB 同步清除
    r = httpx.delete(f"{BASE}/api/admin/knowledge/{doc['id']}", headers=AH)
    assert r.status_code == 200
    from app.rag import vectorstore
    remain = vectorstore.count_chunks()
    assert remain == 0, f"删除后 ChromaDB 应为空, 实际 {remain}"
    print(f"✅ 删除知识文档，ChromaDB 同步清空 (剩余 chunks={remain})")

    print("\n🎉 Phase 8 管理后台端到端通过")


loop = asyncio.get_event_loop()
loop.run_until_complete(main())
