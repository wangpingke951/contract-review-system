"""生成一份"高风险"测试合同 —— 用于完整验证审查链路的 9 个维度

刻意埋设的风险点（便于肉眼核对审查报告是否命中）：
- 条款缺失：无不可抗力 / 无数据合规与个人信息保护 / 无质保与维护 / 无附件清单
- 金额异常：合同总额被"以实际验收为准"架空、甲方单方调减权、180 天账期、背靠背付款
- 期限风险：起算日依赖甲方通知、自动续约无异议期与通知期
- 主体资格：乙方为"工作室"，缺统一社会信用代码 / 住所地 / 法定代表人
- 违约责任：乙方 1%/日上不封顶，甲方逾期付款完全免责，解除权不对等
- 知识产权：乙方背景知识产权被一并收走，登记费用转嫁
- 保密条款：无限期保密、200 万违约金（远超 90 万合同额）、义务单向
- 争议解决：管辖地与法律适用均单方有利
- 正面发现：第八条廉洁条款双向对等、期限合理、救济明确

运行：
    python backend/tests/make_risky_contract.py
输出：
    项目根目录 test_contract_risky.docx
"""
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from helpers import _asset  # noqa: E402
from docx import Document  # noqa: E402

doc = Document()

doc.add_heading("软件开发与技术服务合同", level=0)

doc.add_paragraph("合同编号：XC-2026-TECH-0187")
doc.add_paragraph("签订地点：成都市高新区")
doc.add_paragraph("签订日期：2026 年 10 月 15 日")

doc.add_paragraph("甲方（委托方）：星辰科技（成都）有限公司")
doc.add_paragraph("乙方（受托方）：老王软件开发工作室")

# ——— 第一条：服务内容 ———
doc.add_heading("第一条 服务内容", level=1)
doc.add_heading("1.1 乙方为甲方开发“智慧供应链管理平台”软件系统，具体功能需求以甲方通过电子邮件发送的《需求说明书》为准，该说明书无需双方签字确认。", level=2)
doc.add_heading("1.2 甲方有权在开发过程中随时调整、追加或删减需求，乙方应无条件配合，不得据此要求增加费用或延长工期。", level=2)
doc.add_heading("1.3 乙方应于甲方通知的交付日完成交付。验收标准与验收流程由甲方自行确定，乙方不得提出异议。", level=2)
doc.add_heading("1.4 本合同未约定系统可用性指标、故障响应时限及质保期限。", level=2)

# ——— 第二条：金额与付款 ———
doc.add_heading("第二条 合同金额与付款", level=1)
doc.add_heading("2.1 本合同总金额为人民币玖拾万元整（¥900,000.00），含增值税专用发票。", level=2)
doc.add_heading("2.2 最终结算金额以甲方实际验收确认的工作量为准，甲方有权根据验收情况单方调减合同金额，乙方应予接受。", level=2)
doc.add_heading("2.3 付款方式：乙方完成全部交付并经甲方书面验收合格后 180 日内，甲方一次性支付。甲方付款以其最终客户向其付款为前提。", level=2)
doc.add_heading("2.4 乙方应在收到款项前先行开具全额发票；因发票不合规导致的付款延误由乙方自行承担。", level=2)

# ——— 第三条：期限与续约 ———
doc.add_heading("第三条 服务期限与续约", level=1)
doc.add_heading("3.1 服务期限自甲方书面通知之日起 12 个月。", level=2)
doc.add_heading("3.2 期限届满前，如双方均未提出异议，本合同自动续约 12 个月，续约期间费用按本合同执行。", level=2)
doc.add_heading("3.3 乙方如需终止续约，应提前书面通知甲方，具体提前期限由甲方确定。", level=2)

# ——— 第四条：知识产权 ———
doc.add_heading("第四条 知识产权归属", level=1)
doc.add_heading("4.1 乙方在本合同履行过程中开发形成的全部工作成果，包括但不限于源代码、文档、设计稿、算法模型，其知识产权均归甲方单独所有。", level=2)
doc.add_heading("4.2 乙方在履行本合同前已拥有的既有技术、通用组件及工具库，亦视为本合同工作成果，一并归甲方所有，乙方不得再用于任何其他项目。", level=2)
doc.add_heading("4.3 乙方应无条件配合甲方办理著作权登记、专利申请等手续，相关费用由乙方承担。", level=2)

# ——— 第五条：保密 ———
doc.add_heading("第五条 保密义务", level=1)
doc.add_heading("5.1 乙方对因履行本合同知悉的甲方全部信息负有保密义务，保密期限为本合同终止后永久有效。", level=2)
doc.add_heading("5.2 乙方违反保密义务的，应支付违约金人民币贰佰万元整（¥2,000,000.00），并赔偿甲方因此遭受的全部损失。", level=2)
doc.add_heading("5.3 甲方对乙方提供的技术信息、报价信息不承担保密义务。", level=2)

# ——— 第六条：违约责任 ———
doc.add_heading("第六条 违约责任", level=1)
doc.add_heading("6.1 乙方逾期交付或交付成果未通过验收的，每逾期一日按合同总金额的 1% 支付违约金，累计计算不设上限。", level=2)
doc.add_heading("6.2 乙方交付成果存在质量问题的，应免费返工直至甲方满意；甲方有权同时要求乙方赔偿全部直接损失及间接损失。", level=2)
doc.add_heading("6.3 甲方逾期付款的，不承担违约责任，乙方不得因此暂停服务或主张任何权利。", level=2)
doc.add_heading("6.4 任何一方均有权提前 30 日书面通知解除本合同，但乙方解除合同须经甲方书面同意。", level=2)

# ——— 第七条：争议解决 ———
doc.add_heading("第七条 争议解决", level=1)
doc.add_heading("7.1 因本合同引起的争议，双方应友好协商解决；协商不成的，提交甲方所在地有管辖权的人民法院诉讼解决。", level=2)
doc.add_heading("7.2 本合同的订立、效力、解释及争议解决均适用甲方所在地法律。", level=2)

# ——— 第八条：正面条款 ———
doc.add_heading("第八条 廉洁与反商业贿赂", level=1)
doc.add_heading("8.1 甲乙双方及其工作人员均不得向对方工作人员提供或索取任何形式的不正当利益，包括但不限于现金、礼品、回扣、有价证券。", level=2)
doc.add_heading("8.2 任何一方发现对方存在违反本条情形的，有权书面通知对方纠正；情节严重的，有权解除本合同并要求赔偿损失。", level=2)
doc.add_heading("8.3 本条义务在本合同终止或解除后继续有效 2 年。", level=2)

# ——— 第九条：其他 ———
doc.add_heading("第九条 其他约定", level=1)
doc.add_heading("9.1 本合同未尽事宜，由甲方负责解释。", level=2)
doc.add_heading("9.2 乙方不得在任何宣传材料、案例展示中使用甲方名称、商标，或披露双方合作关系。", level=2)
doc.add_heading("9.3 本合同项下的通知以甲方发出的电子邮件为准，发送至乙方在甲方系统中登记的电子邮箱即视为有效送达。", level=2)
doc.add_heading("9.4 本合同未就不可抗力、数据合规与个人信息保护、合同附件及其优先效力作出约定。", level=2)

doc.add_paragraph("（以下无正文）")
doc.add_paragraph("甲方（盖章）：                    乙方（盖章）：")
doc.add_paragraph("授权代表（签字）：                授权代表（签字）：")

out = _asset("test_contract_risky.docx")
doc.save(out)
print("[OK] generated:", out)
